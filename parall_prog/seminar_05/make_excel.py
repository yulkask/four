# Семинар 5. Таблицы и графики ускорения/эффективности метода Гаусса в Excel
# Запуск из папки parall_prog: python3 seminar_05/make_excel.py  (нужен openpyxl)
import csv
import re

from openpyxl import Workbook
from openpyxl.chart import LineChart, Reference
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

SIZES = [1000, 2000, 4000]
THREADS = [1, 2, 4, 8]

times = {}
with open("results/seminar5_gauss.csv") as f:
    for row in csv.DictReader(f):
        key = "serial" if row["version"] == "serial" else int(row["threads"])
        times[(int(row["n"]), key)] = float(row["time_sec"])

FONT = "Arial"
bold = Font(name=FONT, bold=True)
regular = Font(name=FONT)
title_font = Font(name=FONT, bold=True, size=12)
header_fill = PatternFill("solid", start_color="D9E1F2")
thin = Side(style="thin", color="A6A6A6")
border = Border(left=thin, right=thin, top=thin, bottom=thin)
center = Alignment(horizontal="center", vertical="center")


def header(ws, row, labels):
    for col, text in enumerate(labels, start=1):
        c = ws.cell(row=row, column=col, value=text)
        c.font = bold
        c.fill = header_fill
        c.border = border
        c.alignment = center


def body(cell, fmt=None):
    cell.font = regular
    cell.border = border
    cell.alignment = center
    if fmt:
        cell.number_format = fmt


wb = Workbook()
ws = wb.active
ws.title = "Результаты"
ws.column_dimensions["A"].width = 26
for col in "BCD":
    ws.column_dimensions[col].width = 16

ws["A1"] = "Метод Гаусса: последовательный и параллельный (OpenMP)"
ws["A1"].font = title_font
ws["A2"] = "Apple M4 (4 производительных + 6 энергоэффективных ядер), clang -O2, лучшее из 3 запусков (4000×4000 — 1 запуск)"
ws["A2"].font = Font(name=FONT, italic=True, color="595959")

# Таблица 1. Время, с (данные из results/seminar5_gauss.csv)
ws["A4"] = "Таблица 1. Время решения, с"
ws["A4"].font = bold
header(ws, 5, ["Версия / потоки"] + [f"{n}×{n}" for n in SIZES])
ws.cell(row=6, column=1, value="Последовательная")
body(ws.cell(row=6, column=1))
for j, n in enumerate(SIZES, start=2):
    body(ws.cell(row=6, column=j, value=times[(n, "serial")]), "0.0000")
for i, p in enumerate(THREADS):
    r = 7 + i
    body(ws.cell(row=r, column=1, value=p))
    for j, n in enumerate(SIZES, start=2):
        body(ws.cell(row=r, column=j, value=times[(n, p)]), "0.0000")
time_first, time_last = 7, 6 + len(THREADS)

# Таблица 2. Ускорение S = T_посл / T_p
ws["A12"] = "Таблица 2. Ускорение S = Tпосл / Tp"
ws["A12"].font = bold
header(ws, 13, ["Потоки"] + [f"{n}×{n}" for n in SIZES])
for i, p in enumerate(THREADS):
    r = 14 + i
    body(ws.cell(row=r, column=1, value=p))
    for j in range(2, 2 + len(SIZES)):
        col = ws.cell(row=6, column=j).column_letter
        body(ws.cell(row=r, column=j, value=f"={col}$6/{col}{time_first + i}"), "0.00")
sp_first = 14

# Таблица 3. Эффективность E = S / p
ws["A19"] = "Таблица 3. Эффективность E = S / p"
ws["A19"].font = bold
header(ws, 20, ["Потоки"] + [f"{n}×{n}" for n in SIZES])
for i, p in enumerate(THREADS):
    r = 21 + i
    body(ws.cell(row=r, column=1, value=p))
    for j in range(2, 2 + len(SIZES)):
        col = ws.cell(row=6, column=j).column_letter
        body(ws.cell(row=r, column=j, value=f"={col}{sp_first + i}/$A{r}"), "0.00")
ef_first = 21

ws["A26"] = "Время в таблице 1 — результат программы gauss.cpp; таблицы 2 и 3 считаются формулами."
ws["A26"].font = Font(name=FONT, italic=True, color="595959")


def line_chart(title, y_title, first_row, anchor):
    ch = LineChart()
    ch.title = title
    ch.x_axis.title = "Число потоков"
    ch.y_axis.title = y_title
    ch.height, ch.width = 8, 15
    data = Reference(ws, min_col=2, max_col=1 + len(SIZES), min_row=first_row - 1, max_row=first_row + len(THREADS) - 1)
    cats = Reference(ws, min_col=1, min_row=first_row, max_row=first_row + len(THREADS) - 1)
    ch.add_data(data, titles_from_data=True)
    ch.set_categories(cats)
    for s in ch.series:
        s.marker.symbol = "circle"
        s.smooth = False
    ch.x_axis.delete = False
    ch.y_axis.delete = False
    ws.add_chart(ch, anchor)


line_chart("Время решения", "Время, с", time_first, "F4")
line_chart("Ускорение", "S", sp_first, "F21")
line_chart("Эффективность", "E", ef_first, "F38")

# Лист с корнями систем из output.txt
ws2 = wb.create_sheet("Системы")
ws2.column_dimensions["A"].width = 12
ws2.column_dimensions["B"].width = 24
ws2.column_dimensions["C"].width = 24
ws2["A1"] = "Корни систем из «Уравнения.docx»"
ws2["A1"].font = title_font
header(ws2, 3, ["Система", "Последовательный", "Параллельный"])
r = 4
name = None
roots = {}
with open("seminar_05/output.txt") as f:
    for line in f:
        m = re.match(r"Система (\S+)\)", line)
        if m:
            name = m.group(1)
            roots[name] = {}
        m = re.search(r"(последоват\.|параллельный):\s+(\(.*\))", line)
        if m and name:
            roots[name][m.group(1)] = m.group(2)
for name, v in roots.items():
    body(ws2.cell(row=r, column=1, value=name + ")"))
    body(ws2.cell(row=r, column=2, value=v.get("последоват.", "")))
    body(ws2.cell(row=r, column=3, value=v.get("параллельный", "")))
    r += 1

wb.calculation.fullCalcOnLoad = True
wb.save("results/seminar5_gauss.xlsx")
print("Сохранено: results/seminar5_gauss.xlsx")
